"""Загрузка документов: файл → текст → чанки → эмбеддинги → Chroma.

Поддерживаются PDF, DOCX, XLSX, CSV, TXT, MD.

CLI (бэкенд при этом должен быть остановлен — Chroma не любит два процесса):
    uv run python -m app.rag.ingest data/demo data/training
Когда бэкенд запущен, загружайте через API: make ingest
"""

import argparse
import csv
import hashlib
import logging
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import get_settings
from app.llm.client import get_llm
from app.log import setup_logging
from app.rag.store import get_collection

log = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}


def read_text(path: Path) -> str:
    match path.suffix.lower():
        case ".pdf":
            return "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        case ".docx":
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                parts.append(rows_to_text(rows))
            return "\n\n".join(parts)
        case ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for ws in wb.worksheets:
                rows = [
                    ["" if v is None else str(v) for v in row]
                    for row in ws.iter_rows(values_only=True)
                ]
                sheets.append(f"Лист «{ws.title}»\n\n{rows_to_text(rows)}")
            return "\n\n".join(sheets)
        case ".csv":
            raw = path.read_text(encoding="utf-8-sig", errors="replace")
            dialect = csv.Sniffer().sniff(raw[:2000], delimiters=",;\t")
            return rows_to_text(list(csv.reader(raw.splitlines(), dialect)))
        case ".txt":
            return path.read_text(encoding="utf-8", errors="replace")
        case ".md":
            text = path.read_text(encoding="utf-8", errors="replace")
            if text.startswith("---\n"):  # шапка учебного модуля (title, required) — не текст
                text = text.removeprefix("---\n").partition("\n---\n")[2]
            return text
    raise ValueError(f"Формат {path.suffix} не поддерживается. Можно: {sorted(SUPPORTED_SUFFIXES)}")


def rows_to_text(rows: list[list[str]]) -> str:
    """Таблица → текст «колонка: значение; ...», по абзацу на строку — чанк понятен без шапки."""
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ""
    header, *body = rows
    lines = [
        "; ".join(
            f"{h.strip()}: {v.strip()}" for h, v in zip(header, row, strict=False) if v.strip()
        )
        for row in body
    ]
    return "\n\n".join(lines)


def split_text(text: str, size: int, overlap: int) -> list[str]:
    """Режем по абзацам и склеиваем их в чанки до size символов.
    Слишком длинный абзац режем окном с перекрытием overlap."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for p in paragraphs:
        if len(p) > size:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(p[i : i + size] for i in range(0, len(p), size - overlap))
            continue
        if current and len(current) + len(p) + 2 > size:
            chunks.append(current)
            current = ""
        current = f"{current}\n\n{p}" if current else p
    if current:
        chunks.append(current)
    return chunks


def ingest_file(path: Path, name: str | None = None, digest: str = "") -> int:
    """Индексирует один файл. Повторная загрузка файла с тем же именем заменяет старые чанки.
    digest — хеш файла из репозитория, по нему autoload_documents() видит, что файл поменялся."""
    s = get_settings()
    name = name or path.name
    chunks = split_text(read_text(path), s.chunk_size, s.chunk_overlap)
    collection = get_collection()
    collection.delete(where={"source": name})
    if not chunks:
        log.warning("%s: текст не извлечён (скан без OCR?)", name)
        return 0
    vectors = get_llm().embed(chunks, kind="doc")
    collection.upsert(
        ids=[f"{name}:{i}" for i in range(len(chunks))],
        embeddings=vectors,
        documents=chunks,
        metadatas=[{"source": name, "chunk": i, "hash": digest} for i in range(len(chunks))],
    )
    log.info("%s: %d чанков", name, len(chunks))
    return len(chunks)


def ingest_path(path: Path) -> dict[str, int]:
    files = [path] if path.is_file() else sorted(p for p in path.rglob("*") if p.is_file())
    return {f.name: ingest_file(f) for f in files if f.suffix.lower() in SUPPORTED_SUFFIXES}


def autoload_documents() -> int:
    """При старте бэкенда загружаем документы фонда из data/demo и data/training, которых ещё нет
    в базе знаний или которые поменялись в репозитории (сверяем по хешу файла). Так новый
    документ из репозитория попадает на сервер сам, без make ingest. Возвращает число новых чанков.
    Ошибка не роняет бэкенд: помощник тогда отвечает по тому, что уже есть, и краткой справке о фонде."""
    try:
        metas = get_collection().get(include=["metadatas"])["metadatas"] or []
        known = {m["source"]: m.get("hash", "") for m in metas}
        data_dir = get_settings().data_dir
        total = 0
        for name in ("demo", "training"):
            folder = data_dir / name
            if not folder.is_dir():
                continue
            for path in sorted(folder.rglob("*")):
                if path.suffix.lower() not in SUPPORTED_SUFFIXES:
                    continue
                digest = file_hash(path)
                if known.get(path.name) != digest:
                    total += ingest_file(path, digest=digest)
        if total:
            log.info("Загрузили новые и изменённые документы фонда: %d чанков", total)
        return total
    except Exception as e:  # сеть, ключи, битый файл — бэкенд всё равно должен подняться
        log.warning("Не удалось загрузить документы фонда при запуске: %s", e)
        return 0


def file_hash(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description="Загрузить документы в Chroma")
    parser.add_argument("paths", nargs="+", type=Path, help="файлы или папки")
    args = parser.parse_args()
    setup_logging(get_settings().log_level)
    total = 0
    for path in args.paths:
        total += sum(ingest_path(path).values())
    log.info("Готово: %d чанков в коллекции «%s»", total, get_collection().name)


if __name__ == "__main__":
    main()
